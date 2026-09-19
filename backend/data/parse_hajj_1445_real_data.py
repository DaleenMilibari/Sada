"""One-off parser for the real SBA Hajj season 1445H media statistics.

Reads backend/data/raw/hajj_1445_media_stats.xlsx (an open-data export from
SBA covering Hajj season 1445H) and writes hajj_1445_platform_totals.json.

Unlike engagement_records.json (synthetic, per-post), this source is a
season-level totals table: one row per platform, with different metrics
populated per platform depending on what each platform's own analytics
expose (e.g. YouTube has no reported "engagement" figure, Meta has no
reported "views" figure). Missing cells are kept as null rather than
zero-filled or guessed - see behavior_analysis_agent.analyze_real_hajj_platforms
for how patterns only compare platforms that actually report a given metric.

Run with `python -m backend.data.parse_hajj_1445_real_data` from the repo root.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

DATA_DIR = Path(__file__).parent
RAW_PATH = DATA_DIR / "raw" / "hajj_1445_media_stats.xlsx"
OUT_PATH = DATA_DIR / "hajj_1445_platform_totals.json"

# Arabic header labels (as they appear in the sheet) -> our metric keys.
HEADER_MAP = {
    "مرات الظهور": "impressions",
    "المشاهدات": "views",
    "التفاعل": "engagement",
    "وقت المشاهدة": "watch_time",
    "مشاهدات الملف الشخصي": "profile_views",
    "الإعجاب": "likes",
}

# Arabic platform row labels -> our platform keys.
PLATFORM_MAP = {
    "منصة x": "x",
    "يوتيوب": "youtube",
    "منصة ميتا": "meta",
    "تيك توك": "tiktok",
}

TOTAL_ROW_LABEL = "اجمالي"


def _parse_number(raw: object) -> float | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    match = re.match(r"^([\d.]+)\s*([KMkm]?)$", text)
    if not match:
        return None
    value, suffix = match.groups()
    value = float(value)
    if suffix.upper() == "K":
        value *= 1_000
    elif suffix.upper() == "M":
        value *= 1_000_000
    return value


def parse(raw_path: Path = RAW_PATH) -> dict:
    wb = openpyxl.load_workbook(raw_path, data_only=True)
    ws = wb["Sheet1"]

    header_row = 4
    headers: dict[int, str] = {}
    for cell in ws[header_row]:
        if cell.value is None:
            continue
        label = str(cell.value).strip()
        metric = HEADER_MAP.get(label)
        if metric:
            headers[cell.column] = metric

    platforms: dict[str, dict] = {}
    total_reported: dict[str, float] = {}

    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        label_cell = row[0] if row[0].column == 3 else next((c for c in row if c.column == 3), None)
        if label_cell is None or label_cell.value is None:
            continue
        label = str(label_cell.value).strip()

        values: dict[str, float | None] = {metric: None for metric in HEADER_MAP.values()}
        for cell in row:
            metric = headers.get(cell.column)
            if metric is None:
                continue
            values[metric] = _parse_number(cell.value)

        if label == TOTAL_ROW_LABEL:
            total_reported = {k: v for k, v in values.items() if v is not None}
            continue

        platform_key = PLATFORM_MAP.get(label)
        if platform_key is None:
            continue
        platforms[platform_key] = {"label_ar": label, **values}

    return {
        "source": {
            "title_ar": "إحصائيات الاعلام لموسم حج 1445هـ",
            "raw_file": "raw/hajj_1445_media_stats.xlsx",
            "note": (
                "Real SBA open data, season-level totals per platform for Hajj "
                "1445H. Not per-post records - each platform reports a "
                "different subset of metrics, so comparisons must only use "
                "platforms that report the same metric (see "
                "analyze_real_hajj_platforms)."
            ),
        },
        "metrics": sorted(set(HEADER_MAP.values())),
        "platforms": platforms,
        "total_reported": total_reported,
    }


def main() -> None:
    result = parse()
    OUT_PATH.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    n_platforms = len(result["platforms"])
    print(f"Wrote {OUT_PATH.name} with {n_platforms} platform(s).")


if __name__ == "__main__":
    main()
